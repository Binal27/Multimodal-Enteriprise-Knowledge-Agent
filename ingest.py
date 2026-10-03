"""
ingest.py — Parses every PDF and PPTX in data/docs/, extracts text and
tables, chunks it, and builds a simple TF-IDF retrieval index.

This uses TF-IDF (keyword/lexical search) rather than dense embeddings on
purpose for this first prototype: it needs no downloaded model and no API
key, so it's fast to set up and fully explainable. Swapping this for dense
embeddings (sentence-transformers or an embedding API) is a natural next
step and is called out in README.md.

Run:  python ingest.py
"""
import os
import pickle
import glob

import pdfplumber
from pptx import Presentation
from sklearn.feature_extraction.text import TfidfVectorizer

DATA_DIR = "data"
INDEX_DIR = "data/index"
CHUNK_SIZE = 500       # characters per chunk
CHUNK_OVERLAP = 80     # overlap between consecutive chunks


def extract_pdf(path):
    """Extract text page by page from a PDF, tagging each chunk with its source page.
    Tables are extracted separately and appended as pipe-delimited rows so
    numbers don't get scrambled by extract_text()'s layout guessing."""
    records = []
    with pdfplumber.open(path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            tables = page.extract_tables()
            table_text = ""
            for table in tables:
                for row in table:
                    row_text = " | ".join(cell or "" for cell in row)
                    table_text += row_text + "\n"
            full_text = (text + "\n" + table_text).strip()
            if full_text:
                records.append({
                    "text": full_text,
                    "source": os.path.basename(path),
                    "location": f"page {page_num}",
                })
    return records


def extract_pptx(path):
    """Extract text from every slide, including any tables on the slide."""
    records = []
    prs = Presentation(path)
    for slide_num, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        texts.append(line.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells)
                    texts.append(row_text)
        slide_text = "\n".join(texts).strip()
        if slide_text:
            records.append({
                "text": slide_text,
                "source": os.path.basename(path),
                "location": f"slide {slide_num}",
            })
    return records


def chunk_text(text, size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    """Simple fixed-size character chunking with overlap."""
    chunks = []
    start = 0
    while start < len(text):
        end = start + size
        chunks.append(text[start:end])
        start = end - overlap
    return chunks


def build_index():
    os.makedirs(INDEX_DIR, exist_ok=True)
    all_chunks = []
    all_metadata = []

    docs_dir = os.path.join(DATA_DIR, "docs")
    paths = glob.glob(os.path.join(docs_dir, "*.pdf")) + \
            glob.glob(os.path.join(docs_dir, "*.pptx"))

    if not paths:
        print(f"No PDF or PPTX files found in {docs_dir}/. Add documents and re-run.")
        return

    for path in paths:
        print(f"Parsing {path} ...")
        if path.lower().endswith(".pdf"):
            records = extract_pdf(path)
        else:
            records = extract_pptx(path)

        for rec in records:
            for chunk in chunk_text(rec["text"]):
                if chunk.strip():
                    all_chunks.append(chunk)
                    all_metadata.append({
                        "source": rec["source"],
                        "location": rec["location"],
                    })

    print(f"Built {len(all_chunks)} chunks from {len(paths)} documents.")

    vectorizer = TfidfVectorizer(stop_words="english")
    matrix = vectorizer.fit_transform(all_chunks)

    with open(os.path.join(INDEX_DIR, "index.pkl"), "wb") as f:
        pickle.dump({
            "vectorizer": vectorizer,
            "matrix": matrix,
            "chunks": all_chunks,
            "metadata": all_metadata,
        }, f)

    print(f"Index saved to {INDEX_DIR}/index.pkl")


if __name__ == "__main__":
    build_index()